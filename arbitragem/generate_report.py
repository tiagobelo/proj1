"""
MVP-03: gera um relatório em XLSX com a calculadora de arbitragem —
lucro estimado, margem, ROI e uma classificação de oportunidade — a
partir do CSV produzido por compare_mercadolivre.py.

Não usa banco de dados: o pipeline inteiro roda em cima de arquivos
CSV/XLSX (amazon_scrape_to_csv.py -> compare_mercadolivre.py -> generate_report.py).

IMPORTANTE sobre as taxas do Mercado Livre: a partir de mar/2026 o custo
operacional passou a ser calculado por peso/dimensão do produto em vez de
uma taxa fixa simples, e a comissão varia por categoria (10-14% no
anúncio clássico, 15-19% no premium). Como este pipeline não coleta
peso/dimensão dos produtos, os valores usados aqui são ESTIMATIVAS —
confirme os valores reais no Simulador de Custos oficial (dentro do
Seller Center) antes de decidir uma compra:
https://www.mercadolivre.com.br/ajuda/formas-de-pagamento/custos

Comissão automática por categoria: se --commission-pct não for passado,
a comissão é escolhida a partir da coluna ml_domain_id do CSV (a
categoria oficial do Mercado Livre, capturada pelo compare_mercadolivre.py
via GeckoAPI) usando o mapa em config/ml_fees.yaml. Passar
--commission-pct explicitamente ignora esse mapa e usa um valor único
fixo para todas as linhas, como antes.

Custo fixo automático por faixa de preço: se --fixed-fee não for
passado, o custo fixo é escolhido pela faixa de preço de venda (o mesmo
--price-basis usado no lucro) em config/ml_fixed_fee_tiers.yaml. Esse
YAML vem vazio (sem faixas confirmadas) — enquanto isso, o script usa
R$ 6,00 fixo para todas as linhas, igual ao comportamento anterior.
Preencha as faixas com os valores do Simulador de Custos para um cálculo
mais preciso. Passar --fixed-fee explicitamente ignora o YAML e usa um
valor único fixo para todas as linhas.

Dependência: openpyxl, PyYAML (pip install openpyxl pyyaml)

Uso:
    python generate_report.py comparacao.csv
    python generate_report.py comparacao.csv --output oportunidades.xlsx
    python generate_report.py comparacao.csv --commission-pct 12 --fixed-fee 6 --shipping-cost 18 --tax-pct 4
    python generate_report.py comparacao.csv --price-basis avg --min-margin-pct 20 --min-roi-pct 25
"""

from __future__ import annotations

import argparse
import csv
from pathlib import Path

import yaml
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

DEFAULT_FEES_CONFIG_PATH = Path(__file__).resolve().parent / "config" / "ml_fees.yaml"
DEFAULT_FIXED_FEE_TIERS_PATH = Path(__file__).resolve().parent / "config" / "ml_fixed_fee_tiers.yaml"
FALLBACK_FIXED_FEE = 6.0

CURRENCY_FORMAT = '"R$" #,##0.00'
PERCENT_FORMAT = "0.0%"

HEADER_FILL = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)

STATUS_FILLS = {
    "🟢": PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid"),
    "🟡": PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid"),
    "🔴": PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid"),
    "⚪": PatternFill(start_color="E5E7EB", end_color="E5E7EB", fill_type="solid"),
}

COLUMNS = [
    ("Produto (Amazon)", 45),
    ("ASIN", 12),
    ("Preço Amazon (custo)", 16),
    ("Preço Mercado Livre", 16),
    ("Base do preço ML", 14),
    ("Nº anúncios ML", 12),
    ("Score do match", 12),
    ("Categoria ML", 22),
    ("Comissão (%)", 12),
    ("Comissão ML", 14),
    ("Custo fixo", 12),
    ("Frete", 12),
    ("Imposto", 12),
    ("Lucro líquido", 14),
    ("Margem", 10),
    ("ROI", 10),
    ("Classificação", 24),
    ("Link Amazon", 40),
    ("Link Mercado Livre", 40),
]


def load_fees_config(path: Path) -> dict:
    with open(path, encoding="utf-8") as f:
        return yaml.safe_load(f)


def load_fixed_fee_tiers(path: Path) -> list[dict]:
    """Carrega as faixas de preço -> custo fixo de config/ml_fixed_fee_tiers.yaml.

    O arquivo pode não existir ou vir com `tiers: []` — nesse caso a lista
    volta vazia e resolve_fixed_fee() cai no valor único FALLBACK_FIXED_FEE.
    """
    if not path.exists():
        return []
    with open(path, encoding="utf-8") as f:
        data = yaml.safe_load(f) or {}
    tiers = data.get("tiers") or []
    # Ordena por up_to crescente, com up_to: null (sem teto) sempre por
    # último — assim a primeira faixa que "cobre" o preço de venda é a
    # correta, sem exigir que o YAML já venha ordenado.
    return sorted(tiers, key=lambda t: (t.get("up_to") is None, t.get("up_to")))


def resolve_fixed_fee(sale_price: float | None, tiers: list[dict], override: float | None) -> float:
    """Se o usuário passou --fixed-fee, esse valor vale para todas as
    linhas (comportamento antigo, ignora config/ml_fixed_fee_tiers.yaml).
    Senão, escolhe pela faixa de preço de venda no YAML. Sem faixas
    configuradas (arquivo vazio) ou sem preço de venda, usa
    FALLBACK_FIXED_FEE."""
    if override is not None:
        return override
    if sale_price is not None:
        for tier in tiers:
            up_to = tier.get("up_to")
            if up_to is None or sale_price <= float(up_to):
                return float(tier["fee"])
    return FALLBACK_FIXED_FEE


def resolve_commission_pct(domain_id: str, fees_config: dict, override: float | None) -> float:
    """Se o usuário passou --commission-pct, esse valor vale para todas as
    linhas (comportamento antigo). Senão, escolhe pelo domainId da
    categoria (ml_domain_id no CSV) usando config/ml_fees.yaml, com
    fallback pro valor 'default' do arquivo quando o domínio não está
    mapeado ou vem vazio."""
    if override is not None:
        return override
    domains = fees_config.get("domains", {})
    if domain_id and domain_id in domains:
        return float(domains[domain_id])
    return float(fees_config.get("default", 12.0))


def parse_float(value: str) -> float | None:
    if value is None or value == "":
        return None
    try:
        return float(value)
    except ValueError:
        return None


def pick_sale_price(row: dict, basis: str) -> float | None:
    key = {"median": "ml_median_price", "avg": "ml_avg_price", "min": "ml_min_price"}[basis]
    return parse_float(row.get(key, ""))


def classify(net_profit: float | None, margin_pct: float | None, roi_pct: float | None, min_margin_pct: float, min_roi_pct: float) -> tuple[str, str]:
    if net_profit is None:
        return "⚪", "Sem dados suficientes"
    if net_profit <= 0:
        return "🔴", "Não recomendado (prejuízo)"
    if margin_pct >= min_margin_pct and roi_pct >= min_roi_pct:
        return "🟢", "Excelente oportunidade"
    return "🟡", "Oportunidade moderada"


def calculate_opportunity(
    row: dict,
    basis: str,
    commission_pct: float,
    fixed_fee_tiers: list[dict],
    fixed_fee_override: float | None,
    shipping_cost: float,
    tax_pct: float,
    min_margin_pct: float,
    min_roi_pct: float,
) -> dict:
    cost = parse_float(row.get("amazon_price", ""))
    sale_price = pick_sale_price(row, basis)
    listing_count = row.get("ml_listing_count", "")
    match_score = row.get("ml_best_match_score", "")
    domain_id = row.get("ml_domain_id", "")

    result = {
        "title": row.get("amazon_title", ""),
        "asin": row.get("asin", ""),
        "cost": cost,
        "sale_price": sale_price,
        "basis": basis,
        "listing_count": listing_count,
        "match_score": match_score,
        "domain_id": domain_id,
        "commission_pct": commission_pct,
        "commission": None,
        "fixed_fee": None,
        "shipping": None,
        "tax": None,
        "net_profit": None,
        "margin_pct": None,
        "roi_pct": None,
        "status_emoji": "⚪",
        "status_label": "Sem dados suficientes",
        "amazon_url": row.get("amazon_url", ""),
        "ml_url": row.get("ml_best_match_url", ""),
    }

    if cost is None or sale_price is None or cost <= 0 or sale_price <= 0:
        return result

    fixed_fee = resolve_fixed_fee(sale_price, fixed_fee_tiers, fixed_fee_override)
    commission = sale_price * commission_pct / 100
    tax = sale_price * tax_pct / 100
    net_profit = sale_price - cost - commission - fixed_fee - shipping_cost - tax
    margin_pct = net_profit / sale_price * 100
    roi_pct = net_profit / cost * 100

    status_emoji, status_label = classify(net_profit, margin_pct, roi_pct, min_margin_pct, min_roi_pct)

    result.update(
        commission=commission,
        fixed_fee=fixed_fee,
        shipping=shipping_cost,
        tax=tax,
        net_profit=net_profit,
        margin_pct=margin_pct,
        roi_pct=roi_pct,
        status_emoji=status_emoji,
        status_label=status_label,
    )
    return result


def build_workbook(opportunities: list[dict]) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "Oportunidades"

    for col_idx, (title, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=title)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.freeze_panes = "A2"

    basis_label = {"median": "Mediana", "avg": "Média", "min": "Mínimo"}

    for row_idx, opp in enumerate(opportunities, start=2):
        values = [
            opp["title"],
            opp["asin"],
            opp["cost"],
            opp["sale_price"],
            basis_label.get(opp["basis"], opp["basis"]),
            opp["listing_count"],
            opp["match_score"],
            opp["domain_id"],
            (opp["commission_pct"] / 100) if opp.get("commission_pct") is not None else None,
            opp["commission"],
            opp["fixed_fee"],
            opp["shipping"],
            opp["tax"],
            opp["net_profit"],
            (opp["margin_pct"] / 100) if opp["margin_pct"] is not None else None,
            (opp["roi_pct"] / 100) if opp["roi_pct"] is not None else None,
            f"{opp['status_emoji']} {opp['status_label']}",
            opp["amazon_url"],
            opp["ml_url"],
        ]
        for col_idx, value in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.fill = STATUS_FILLS[opp["status_emoji"]]
            header = COLUMNS[col_idx - 1][0]
            if header in ("Preço Amazon (custo)", "Preço Mercado Livre", "Comissão ML", "Custo fixo", "Frete", "Imposto", "Lucro líquido"):
                cell.number_format = CURRENCY_FORMAT
            elif header in ("Margem", "ROI", "Comissão (%)"):
                cell.number_format = PERCENT_FORMAT

    return wb


def main() -> None:
    parser = argparse.ArgumentParser(description="Gera relatório XLSX de oportunidades de arbitragem Amazon -> Mercado Livre")
    parser.add_argument("input_csv", help="CSV gerado por compare_mercadolivre.py")
    parser.add_argument("--output", default="oportunidades.xlsx", help="Caminho do XLSX de saída")
    parser.add_argument("--price-basis", choices=["median", "avg", "min"], default="median", help="Qual preço do Mercado Livre usar como preço de venda esperado (padrão: median, mais resistente a outliers)")
    parser.add_argument("--commission-pct", type=float, default=None, help="Comissão do Mercado Livre em %% fixa para todas as linhas. Se omitido, escolhe automaticamente por categoria (ml_domain_id) usando config/ml_fees.yaml — confirme os valores no Simulador de Custos")
    parser.add_argument("--fees-config", default=str(DEFAULT_FEES_CONFIG_PATH), help="Caminho do YAML de comissão por categoria (padrão: config/ml_fees.yaml)")
    parser.add_argument("--fixed-fee", type=float, default=None, help="Custo fixo em R$ fixo para TODAS as linhas, ignorando config/ml_fixed_fee_tiers.yaml. Se omitido, escolhe automaticamente pela faixa de preço de venda nesse YAML (fallback R$ 6,00 se o YAML estiver vazio) — desde mar/2026 o ML calcula isso por peso/dimensão; confirme no Simulador de Custos")
    parser.add_argument("--fixed-fee-config", default=str(DEFAULT_FIXED_FEE_TIERS_PATH), help="Caminho do YAML de custo fixo por faixa de preço (padrão: config/ml_fixed_fee_tiers.yaml)")
    parser.add_argument("--shipping-cost", type=float, default=0.0, help="Frete estimado em R$ que o vendedor absorve (padrão 0 — normalmente é um custo real relevante, ajuste para o seu caso)")
    parser.add_argument("--tax-pct", type=float, default=0.0, help="Imposto sobre a venda em %% (padrão 0 — depende do seu regime tributário, ex.: Simples Nacional)")
    parser.add_argument("--min-margin-pct", type=float, default=15.0, help="Margem mínima (%%) para classificar como 'excelente oportunidade' (padrão 15)")
    parser.add_argument("--min-roi-pct", type=float, default=20.0, help="ROI mínimo (%%) para classificar como 'excelente oportunidade' (padrão 20)")
    args = parser.parse_args()

    if args.shipping_cost == 0.0:
        print("Aviso: --shipping-cost está em 0. Frete costuma ser um custo real e relevante — considere ajustar.\n")

    fees_config = load_fees_config(Path(args.fees_config))
    fixed_fee_tiers = load_fixed_fee_tiers(Path(args.fixed_fee_config))
    if args.fixed_fee is None and not fixed_fee_tiers:
        print(
            f"Aviso: {args.fixed_fee_config} não tem faixas configuradas (tiers: []) — "
            f"usando R$ {FALLBACK_FIXED_FEE:.2f} fixo para todas as linhas. Preencha o "
            "YAML com os valores do Simulador de Custos para um cálculo por faixa de preço.\n"
        )

    with open(args.input_csv, encoding="utf-8-sig") as f:
        rows = list(csv.DictReader(f))

    opportunities = [
        calculate_opportunity(
            row,
            args.price_basis,
            resolve_commission_pct(row.get("ml_domain_id", ""), fees_config, args.commission_pct),
            fixed_fee_tiers,
            args.fixed_fee,
            args.shipping_cost,
            args.tax_pct,
            args.min_margin_pct,
            args.min_roi_pct,
        )
        for row in rows
    ]

    wb = build_workbook(opportunities)
    wb.save(args.output)

    counts: dict[str, int] = {}
    for opp in opportunities:
        counts[opp["status_emoji"]] = counts.get(opp["status_emoji"], 0) + 1
    print(f"Relatório salvo em {args.output}")
    print(
        f"  🟢 {counts.get('🟢', 0)} excelente(s)  "
        f"🟡 {counts.get('🟡', 0)} moderada(s)  "
        f"🔴 {counts.get('🔴', 0)} não recomendada(s)  "
        f"⚪ {counts.get('⚪', 0)} sem dados"
    )


if __name__ == "__main__":
    main()
